import os
import re
import sys
import json
import argparse
import difflib
import sqlite3
import pandas as pd
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from datetime import datetime, date, timedelta
from pathlib import Path
import openpyxl
import urllib3
import time
import warnings
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock

# --- ИМПОРТЫ ДЛЯ MAX ---
# Убедитесь, что папка max_news лежит рядом со скриптом
from playwright.sync_api import sync_playwright
from max_news.api_collector import ApiCollector, MOSCOW, date_bounds
from max_news.auth import ensure_max_ready
from max_news.protocol import MaxProtocol, SOCKET_HOOK

# --- ИМПОРТЫ ДЛЯ TELEGRAM/VK (та же схема входа через постоянный профиль, что и в MAX) ---
from news_runtime.social_browser import AuthenticationRequired, SocialBrowser, SocialBrowserError

# --- ИМПОРТ ДЛЯ GIGACHAT (авторизация с защёлкой вместо ручного OAuth в этом файле) ---
from news_runtime.gigachat import GigaChatAuthError, GigaChatClient, GigaChatError

# Отключаем системные предупреждения
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
os.system("") 

# --- НАСТРОЙКИ ПУТЕЙ И API ---
BASE_DIR = Path(__file__).resolve().parent
SOURCES_FILE = BASE_DIR / "Мониторинг_Источники.xlsx"
NEWS_TEMPLATE_FILE = BASE_DIR / "Мониторинг_Новости.xlsx"
NEWS_OUTPUT_FILE = BASE_DIR / "Мониторинг_Новости_заполненный.xlsx"
DB_FILE = BASE_DIR / "news_database.db"
MAX_PROFILE_DIR = BASE_DIR / ".data" / "max-profile"
TELEGRAM_PROFILE_DIR = BASE_DIR / ".data" / "telegram-profile"
PATHS_CONFIG_FILE = BASE_DIR / "mega_parser_paths.json"  # пути, сохранённые из меню (не сам ключ/БД-данные)

DEFAULT_SOURCES_FILE = SOURCES_FILE
DEFAULT_NEWS_TEMPLATE_FILE = NEWS_TEMPLATE_FILE
DEFAULT_NEWS_OUTPUT_FILE = NEWS_OUTPUT_FILE
DEFAULT_DB_FILE = DB_FILE

def load_paths_config():
    """Подтягивает пути, сохранённые через меню «Пути к файлам» в прошлый раз.
    Порядок приоритета: CLI-флаги > сохранённый конфиг > пути по умолчанию."""
    global SOURCES_FILE, NEWS_TEMPLATE_FILE, NEWS_OUTPUT_FILE, DB_FILE
    if not PATHS_CONFIG_FILE.exists():
        return
    try:
        data = json.loads(PATHS_CONFIG_FILE.read_text(encoding="utf-8"))
    except (ValueError, OSError) as exc:
        logger.warning(f"⚠️ Не удалось прочитать {PATHS_CONFIG_FILE.name}: {exc}. Использую пути по умолчанию.")
        return
    if data.get("sources"): SOURCES_FILE = Path(data["sources"])
    if data.get("news_template"): NEWS_TEMPLATE_FILE = Path(data["news_template"])
    if data.get("output"): NEWS_OUTPUT_FILE = Path(data["output"])
    if data.get("db"): DB_FILE = Path(data["db"])

def save_paths_config():
    data = {
        "sources": str(SOURCES_FILE),
        "news_template": str(NEWS_TEMPLATE_FILE),
        "output": str(NEWS_OUTPUT_FILE),
        "db": str(DB_FILE),
    }
    PATHS_CONFIG_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

def _load_gigachat_credentials() -> str:
    """Ключ берётся из переменной окружения или локального (не в git) файла —
    не хранится в исходниках. См. gigachat_credentials.example.txt."""
    env_value = os.environ.get("GIGACHAT_CREDENTIALS", "").strip()
    if env_value:
        return env_value
    creds_file = BASE_DIR / "gigachat_credentials.txt"
    if creds_file.exists():
        file_value = creds_file.read_text(encoding="utf-8").strip()
        if file_value:
            return file_value
    raise RuntimeError(
        "Не задан ключ GigaChat. Укажите переменную окружения GIGACHAT_CREDENTIALS "
        "или создайте файл gigachat_credentials.txt рядом со скриптом "
        "(скопируйте gigachat_credentials.example.txt и впишите свой ключ)."
    )

GIGACHAT_CREDENTIALS = _load_gigachat_credentials()
START_DATE = date(2026, 7, 1)  # нижняя граница проекта — раньше не собираем, даже при первом запуске
MAX_WORKERS = 8  # Потоки для обычных сайтов (I/O-bound: сеть, а не CPU — можно много)

# Инкрементальное окно: вместо того чтобы каждый раз пересобирать всё от START_DATE,
# фактическое начало окна — дата предыдущего успешного запуска (с нахлёстом на случай
# пропусков), но не раньше START_DATE. См. compute_effective_start_date(). Через меню
# «Период поиска новостей» можно один раз задать окно вручную — действует только на
# ближайший запуск, дальше снова автоматический режим (см. run()).
INCREMENTAL_OVERLAP_DAYS = 2
EFFECTIVE_START_DATE = START_DATE  # переопределяется в run() перед сбором
EFFECTIVE_END_DATE = date.today()  # переопределяется в run() перед сбором

# Пауза между запросами к GigaChat не зависит от MAX_WORKERS: сколько бы потоков
# ни нашли статьи одновременно, на сам GigaChat они пойдут не чаще этого интервала.
# Иначе рост MAX_WORKERS напрямую увеличивал бы нагрузку/риск 429 от GigaChat.
GIGACHAT_MIN_INTERVAL = 0.6  # секунд между запросами к GigaChat

# После стольких провалов подряд источник уходит в чёрный список и пропускается,
# пока его не сбросят: `python mega_parser.py --reset-blacklist` (или --list-blacklist).
SOURCE_FAILURE_THRESHOLD = 5

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# --- НАСТРОЙКА ЛОГИРОВАНИЯ И СТАТИСТИКИ ---
class ParserStats:
    def __init__(self):
        self.lock = Lock()
        self._reset_locked()

    def _reset_locked(self):
        self.web_sources = 0
        self.max_sources = 0
        self.telegram_sources = 0
        self.blacklisted_skipped = 0
        self.links_found = 0
        self.ai_rejected = 0
        self.gigachat_calls = 0
        self.saved = 0
        self.errors = 0

    def reset(self):
        """Вызывается в начале run() — иначе при нескольких запусках подряд
        в одном сеансе меню сводка показывала бы сумму всех запусков, а не
        только текущего."""
        with self.lock:
            self._reset_locked()

    def add(self, metric, count=1):
        with self.lock:
            setattr(self, metric, getattr(self, metric) + count)

stats = ParserStats()
db_lock = Lock()

class ColoredConsoleFormatter(logging.Formatter):
    def format(self, record):
        msg = record.getMessage()
        time_str = self.formatTime(record, "%H:%M:%S")
        GRAY, RESET = "\033[90m", "\033[0m"
        GREEN, YELLOW, RED = "\033[92m", "\033[93m", "\033[91m"
        CYAN, MAGENTA, WHITE = "\033[96m", "\033[95m", "\033[97m"
        
        if "✅" in msg: color = GREEN
        elif "🗑️" in msg or "Пропуск" in msg: color = YELLOW
        elif "❌" in msg or "Ошибка" in msg: color = RED
        elif "🤖" in msg or "✨" in msg: color = CYAN
        elif "Парсинг WEB" in msg: color = MAGENTA
        elif "Парсинг MAX" in msg: color = MAGENTA
        elif "Парсинг Telegram" in msg: color = MAGENTA
        else: color = WHITE
        
        return f"{GRAY}[{time_str}]{RESET} {color}{msg}{RESET}"

logger = logging.getLogger()
logger.setLevel(logging.INFO)
file_handler = logging.FileHandler("parser.log", encoding="utf-8")
file_handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
logger.addHandler(file_handler)
console_handler = logging.StreamHandler()
console_handler.setFormatter(ColoredConsoleFormatter())
logger.addHandler(console_handler)

# --- 1. БАЗА ДАННЫХ И GIGACHAT (ОБЩЕЕ ЯДРО) ---
def init_db():
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute('''CREATE TABLE IF NOT EXISTS parsed_news (
            url TEXT PRIMARY KEY, title TEXT, source TEXT, municipality TEXT, pub_date TEXT, category TEXT, stage TEXT, summary TEXT, facts TEXT, persons TEXT, organizations TEXT, inn TEXT, location TEXT, created_at TEXT)''')
        # Миграция для БД, созданных до появления разбивки по движкам —
        # старые строки останутся с platform=NULL (эта информация нигде не
        # логировалась раньше, восстановить её задним числом нельзя).
        existing_columns = {row[1] for row in cursor.execute("PRAGMA table_info(parsed_news)")}
        if "platform" not in existing_columns:
            cursor.execute("ALTER TABLE parsed_news ADD COLUMN platform TEXT")
        cursor.execute('''CREATE TABLE IF NOT EXISTS rejected_urls (
            url TEXT PRIMARY KEY, reason TEXT, created_at TEXT)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS source_health (
            source_url TEXT PRIMARY KEY, source_name TEXT, consecutive_failures INTEGER NOT NULL DEFAULT 0,
            last_failed_at TEXT, blacklisted_at TEXT)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS run_state (
            key TEXT PRIMARY KEY, value TEXT)''')
        conn.commit()
        conn.close()

def get_last_run_date():
    return _get_run_state_date("last_run_date")

def set_last_run_date(value):
    _set_run_state_date("last_run_date", value)

def _get_run_state_date(key):
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        row = conn.execute("SELECT value FROM run_state WHERE key = ?", (key,)).fetchone()
        conn.close()
    if not row or not row[0]:
        return None
    try:
        return datetime.strptime(row[0], "%Y-%m-%d").date()
    except ValueError:
        return None

def _set_run_state_date(key, value):
    """value=None удаляет ключ (возврат к автоматическому режиму)."""
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        if value is None:
            conn.execute("DELETE FROM run_state WHERE key = ?", (key,))
        else:
            conn.execute(
                "INSERT INTO run_state (key, value) VALUES (?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (key, value.strftime("%Y-%m-%d")),
            )
        conn.commit()
        conn.close()

def get_manual_start_date():
    return _get_run_state_date("manual_start_date")

def set_manual_start_date(value):
    _set_run_state_date("manual_start_date", value)

def get_manual_end_date():
    return _get_run_state_date("manual_end_date")

def set_manual_end_date(value):
    _set_run_state_date("manual_end_date", value)

def compute_effective_start_date():
    """Ручная дата (задана через меню) имеет приоритет и действует один раз —
    run() сбрасывает её обратно в автоматический режим после применения.
    Без неё: первый запуск — с START_DATE, дальше — с даты предыдущего запуска
    минус нахлёст (не теряем публикации у границы), но никогда раньше START_DATE."""
    manual = get_manual_start_date()
    if manual is not None:
        return manual
    last_run = get_last_run_date()
    if last_run is None:
        return START_DATE
    return max(START_DATE, last_run - timedelta(days=INCREMENTAL_OVERLAP_DAYS))

def compute_effective_end_date():
    manual = get_manual_end_date()
    return manual if manual is not None else date.today()

def record_source_result(url, name, success):
    """Считает провалы источника подряд; после SOURCE_FAILURE_THRESHOLD источник
    уходит в чёрный список и пропускается в следующих запусках, пока его не сбросят."""
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        if success:
            conn.execute(
                "INSERT INTO source_health (source_url, source_name, consecutive_failures, last_failed_at, blacklisted_at) "
                "VALUES (?, ?, 0, NULL, NULL) "
                "ON CONFLICT(source_url) DO UPDATE SET source_name=excluded.source_name, consecutive_failures=0, "
                "last_failed_at=NULL, blacklisted_at=NULL",
                (url, name),
            )
        else:
            row = conn.execute(
                "SELECT consecutive_failures, blacklisted_at FROM source_health WHERE source_url = ?", (url,)
            ).fetchone()
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            failures = (row[0] if row else 0) + 1
            blacklisted_at = (row[1] if row else None) or (now if failures >= SOURCE_FAILURE_THRESHOLD else None)
            was_blacklisted = bool(row and row[1])
            conn.execute(
                "INSERT INTO source_health (source_url, source_name, consecutive_failures, last_failed_at, blacklisted_at) "
                "VALUES (?, ?, ?, ?, ?) "
                "ON CONFLICT(source_url) DO UPDATE SET source_name=excluded.source_name, "
                "consecutive_failures=excluded.consecutive_failures, last_failed_at=excluded.last_failed_at, "
                "blacklisted_at=excluded.blacklisted_at",
                (url, name, failures, now, blacklisted_at),
            )
            if blacklisted_at and not was_blacklisted:
                logger.warning(
                    f"🚫 Источник «{name}» в чёрном списке после {failures} провалов подряд: "
                    f"{url[:70]}. Сброс: python mega_parser.py --reset-blacklist \"{name}\""
                )
        conn.commit()
        conn.close()

def get_blacklisted_urls():
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        rows = conn.execute("SELECT source_url FROM source_health WHERE blacklisted_at IS NOT NULL").fetchall()
        conn.close()
    return {r[0] for r in rows}

def reset_blacklist(target=None):
    """target=None сбрасывает весь чёрный список; иначе — по точной ссылке или подстроке названия."""
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        if target:
            cur = conn.execute(
                "DELETE FROM source_health WHERE source_url = ? OR source_name LIKE ?",
                (target, f"%{target}%"),
            )
        else:
            cur = conn.execute("DELETE FROM source_health")
        removed = cur.rowcount
        conn.commit()
        conn.close()
    return removed

def print_blacklist():
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        rows = conn.execute(
            "SELECT source_name, source_url, consecutive_failures, last_failed_at FROM source_health "
            "WHERE blacklisted_at IS NOT NULL ORDER BY last_failed_at DESC"
        ).fetchall()
        conn.close()
    if not rows:
        print("Чёрный список источников пуст.")
        return
    print(f"В чёрном списке источников: {len(rows)}")
    for name, url, failures, last_failed_at in rows:
        print(f"  - {name} ({url}) — {failures} провалов подряд, последний: {last_failed_at}")

def is_url_processed(url):
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        res = conn.execute("SELECT 1 FROM parsed_news WHERE url = ?", (url,)).fetchone() or conn.execute("SELECT 1 FROM rejected_urls WHERE url = ?", (url,)).fetchone()
        conn.close()
        return bool(res)

def save_rejected_url(url, reason):
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        try:
            conn.execute("INSERT INTO rejected_urls (url, reason, created_at) VALUES (?, ?, ?)", (url, reason, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()

def save_item_to_db(item):
    def safe_str(val): return ", ".join(str(x) for x in val) if isinstance(val, list) else (str(val) if val is not None else "")
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        success = False
        try:
            # Именованные колонки, а не позиционный VALUES(...) — platform
            # добавлена ALTER TABLE уже после создания таблицы, позиционная
            # вставка перестала бы совпадать по количеству столбцов.
            conn.execute('''INSERT INTO parsed_news
                (url, title, source, municipality, pub_date, category, stage, summary, facts, persons, organizations, inn, location, created_at, platform)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                         (safe_str(item["url"]), safe_str(item["title"]), safe_str(item["source"]), safe_str(item["municipality"]), safe_str(item["pub_date"]), safe_str(item.get("category", "")), safe_str(item.get("stage", "")), safe_str(item.get("summary", "")), safe_str(item.get("facts", "")), safe_str(item.get("persons", "")), safe_str(item.get("organizations", "")), safe_str(item.get("inn", "")), safe_str(item.get("location", "")), datetime.now().strftime("%Y-%m-%d %H:%M:%S"), safe_str(item.get("platform", ""))))
            conn.commit()
            success = True
        except sqlite3.IntegrityError: pass
        conn.close()
        return success

def is_fuzzy_duplicate(new_title, threshold=0.85):
    """Проверяет, нет ли в базе новости с очень похожим заголовком"""
    if not new_title:
        return False
        
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        # Берем последние 300 сохраненных новостей для быстрой проверки
        cursor.execute("SELECT title FROM parsed_news ORDER BY created_at DESC LIMIT 300")
        existing_titles = [row[0] for row in cursor.fetchall() if row[0]]
        conn.close()
        
    new_title_lower = new_title.lower()
    for old_title in existing_titles:
        # Вычисляем процент схожести двух строк
        similarity = difflib.SequenceMatcher(None, new_title_lower, old_title.lower()).ratio()
        if similarity >= threshold:
            return True
    return False

def _describe_request_error(exc: requests.exceptions.RequestException) -> str:
    """Короткое, читаемое описание сетевой ошибки вместо длинного дампа исключения."""
    if isinstance(exc, requests.exceptions.HTTPError) and exc.response is not None:
        return f"HTTP {exc.response.status_code}"
    if isinstance(exc, requests.exceptions.SSLError):
        return "ошибка проверки TLS-сертификата"
    if isinstance(exc, requests.exceptions.Timeout):
        return "превышено время ожидания ответа"
    if isinstance(exc, requests.exceptions.ConnectionError):
        return "не удалось подключиться (DNS или сеть недоступны)"
    return str(exc)[:150]

def _is_transient_request_error(exc: requests.exceptions.RequestException) -> bool:
    """4xx (кроме 429) — постоянная ошибка источника, повторять его бессмысленно."""
    if isinstance(exc, requests.exceptions.HTTPError) and exc.response is not None:
        status = exc.response.status_code
        return status == 429 or 500 <= status < 600
    return True

def request_with_retry(method, url, max_retries=3, remember_dead_url=None, **kwargs):
    """remember_dead_url: если задан (обычно = url конкретной статьи, а не
    страницы-листинга источника) — при окончательном провале ссылка пишется в
    rejected_urls, чтобы не пытаться её заново на каждом следующем запуске.
    Для листинга источника не передавайте — там свой механизм (чёрный список
    источников, record_source_result), эта ссылка не про отдельную статью."""
    delay = 2
    for attempt in range(max_retries):
        try:
            res = requests.request(method, url, **kwargs)
            res.raise_for_status()
            return res
        except requests.exceptions.RequestException as e:
            if attempt == max_retries - 1 or not _is_transient_request_error(e):
                stats.add("errors")
                description = _describe_request_error(e)
                logger.error(f"❌ Сбой сети {url[:70]}: {description}")
                if remember_dead_url:
                    save_rejected_url(remember_dead_url, f"Сеть: {description}")
                return None
            time.sleep(delay)
            delay *= 2

# GigaChatClient сам обновляет токен и один раз "защёлкивает" отказ авторизации
# (401/403) вместо того, чтобы долбить API на каждую статью одним и тем же ключом.
gigachat_client = GigaChatClient(GIGACHAT_CREDENTIALS, verify=False)
_gigachat_auth_failure_logged = False
_gigachat_auth_lock = Lock()

# Отдельный от MAX_WORKERS ограничитель частоты — см. GIGACHAT_MIN_INTERVAL выше.
_gigachat_rate_lock = Lock()
_gigachat_last_call_at = 0.0

def _throttle_gigachat():
    global _gigachat_last_call_at
    with _gigachat_rate_lock:
        wait = _gigachat_last_call_at + GIGACHAT_MIN_INTERVAL - time.time()
        if wait > 0:
            time.sleep(wait)
        _gigachat_last_call_at = time.time()

def call_gigachat_api(system_prompt, user_message):
    global _gigachat_auth_failure_logged
    if _gigachat_auth_failure_logged:
        return None  # ключ уже отклонён в этом запуске — не тратим время и сеть впустую
    _throttle_gigachat()
    stats.add("gigachat_calls")
    try:
        return gigachat_client.complete_json(system_prompt, user_message)
    except GigaChatAuthError as e:
        stats.add("errors")
        with _gigachat_auth_lock:
            if not _gigachat_auth_failure_logged:
                _gigachat_auth_failure_logged = True
                logger.error(f"❌ GigaChat: {e} Дальнейшие запросы к GigaChat не выполняются в этом запуске.")
        return None
    except GigaChatError as e:
        stats.add("errors")
        logger.error(f"❌ GigaChat: {e}")
        return None

def is_news_relevant(title, text, municipality):
    system_prompt = f"""Ты строгий фильтр новостей. Муниципалитет: {municipality}.
    Правила:
    1. Если новость про концерты, праздники, афиши, советы или рекламу -> is_relevant: false
    2. Если новость про ремонты, стройку, бюджет, ДТП, экономику, решения властей -> is_relevant: true

    ПРИМЕР 1 (МУСОР):
    Текст: Завтра в Доме культуры пройдет концерт хора ветеранов, приглашаем всех желающих.
    Ответ: {{"is_relevant": false, "reason": "Анонс рядового культурного мероприятия без управленческих решений"}}

    ПРИМЕР 2 (ПОЛЕЗНОЕ):
    Текст: Подрядчик завершил укладку асфальта на улице Ленина на неделю раньше срока.
    Ответ: {{"is_relevant": true, "reason": "Завершение значимого инфраструктурного проекта"}}

    Верни ТОЛЬКО JSON: {{"is_relevant": true/false, "reason": "краткая причина"}}"""
    return call_gigachat_api(system_prompt, f"Заголовок: {title}\nТекст: {text}")

def extract_news_details(title, text, municipality):
    system_prompt = f"""Ты помощник по извлечению фактов. Муниципалитет: {municipality}. Заполни поля по тексту новости. 
    Категория (ВЫБЕРИ ОДНУ ИЗ): Власть и кадровые решения, Бюджет и муниципальные финансы, Закупки и контракты, Строительство и инфраструктура, ЖКХ и благоустройство, Экономика, бизнес и занятость, Социальная сфера, Образование, Здравоохранение, Происшествия и безопасность, Экология, Обращения и проблемы жителей, Культура, спорт и туризм, Банки и финансовые сервисы, Другое.
    Стадия (ВЫБЕРИ ОДНУ ИЗ): Анонс, Решение принято, Финансирование выделено, Закупка объявлена, Работы начаты, В процессе, Завершено / открыто, Перенесено / приостановлено, Проблема / происшествие, Результат.

    ПРИМЕР ИДЕАЛЬНОГО ОТВЕТА:
    {{
        "category": "Строительство и инфраструктура",
        "stage": "Завершено / открыто",
        "summary": "Завершено строительство нового детского сада на 150 мест в южном районе.",
        "facts": "150 мест, 50 млн рублей, сдано на 2 месяца раньше",
        "persons": "Иванов И.И. (мэр города)",
        "organizations": "ООО СтройТрест",
        "inn": "1101234567",
        "location": "ул. Мира, 15"
    }}

    Верни ТОЛЬКО JSON: {{"category": "категория", "stage": "стадия", "summary": "1-2 предложения сути", "facts": "цифры, деньги, сроки", "persons": "ФИО", "organizations": "ведомства, компании", "inn": "ИНН", "location": "улица, объект"}}"""
    return call_gigachat_api(system_prompt, f"Заголовок: {title}\nТекст: {text}")

def process_and_save_ai(url, title, text, source_name, municipality, pub_date, platform):
    """Общая функция прогона через ИИ и сохранения для обоих парсеров"""
    
    # Отсекаем перепечатки одних и тех же новостей разными сайтами
    if is_fuzzy_duplicate(title):
        logger.info(f"🟡 [ДУБЛИКАТ]: '{title[:45]}...' уже есть в базе. Пропускаем.")
        return False

    logger.info(f"🤖 AI Проверка: {title[:55]}...")
    filter_data = is_news_relevant(title, text, municipality)
    
    if not filter_data: 
        return False
        
    if not filter_data.get("is_relevant"):
        stats.add("ai_rejected")
        reason = filter_data.get("reason", "Мусор")
        save_rejected_url(url, reason)
        logger.info(f"🗑️ [ОТКЛОНЕНО]: {reason}")
        return False

    logger.info(f"✨ Извлечение данных: {title[:30]}...")
    ai_data = extract_news_details(title, text, municipality) or {}

    item = {
        "url": url, "title": title, "source": source_name, "municipality": municipality,
        "pub_date": pub_date, "summary": text[:500], "platform": platform,
        **{k: ai_data.get(k, "") for k in ["category", "stage", "summary", "facts", "persons", "organizations", "inn", "location"]}
    }
    if save_item_to_db(item):
        stats.add("saved")
        logger.info(f"✅ [СОХРАНЕНО | {item.get('category', 'Б/К')}] {title[:45]}...")
        return True
    return False

# --- 2. ПАРСИНГ ОБЫЧНЫХ САЙТОВ (WEB) ---
def clean_text(text):
    if not text: return ""
    return re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F]', '', str(text))[:32000].strip()

def extract_date_obj(text):
    if not text: return None
    text = str(text).strip().lower()
    match_iso = re.search(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', text)
    if match_iso:
        y, m, d = map(int, match_iso.groups())
        if 2000 <= y <= 2030: return date(y, m, d)
    return None

def parse_web_article(url):
    res = request_with_retry("GET", url, headers=HEADERS, timeout=30, verify=False, remember_dead_url=url)
    if not res or "text/html" not in res.headers.get("Content-Type", "").lower(): return None
    res.encoding = res.apparent_encoding or 'utf-8'
    soup = BeautifulSoup(res.text, "html.parser")
    
    pub_date_obj = None
    for tag in soup.find_all("meta", attrs={"property": "article:published_time"}):
        if tag.get("content"): pub_date_obj = extract_date_obj(tag["content"])
    if not pub_date_obj and soup.find("time"):
        pub_date_obj = extract_date_obj(soup.find("time").get("datetime") or soup.find("time").get_text())

    if not pub_date_obj or not (EFFECTIVE_START_DATE <= pub_date_obj <= EFFECTIVE_END_DATE): return None

    title = soup.find("h1").get_text(" ", strip=True) if soup.find("h1") else ""
    if not title or len(title) < 10:
        if soup.title and soup.title.string: title = soup.title.string.split("—")[0].split("-")[0].strip()

    paragraphs = [p.get_text(" ", strip=True) for p in soup.find_all("p") if len(p.get_text(" ", strip=True)) > 40]
    summary = " ".join(paragraphs[:3]) if paragraphs else ""
    
    return {"title": clean_text(title or summary[:100]+"..."), "pub_date": pub_date_obj.strftime("%Y-%m-%d"), "summary": clean_text(summary or title)}

def process_single_web_news(full_url, source_name, municipality):
    if is_url_processed(full_url): return
    details = parse_web_article(full_url)
    if details:
        process_and_save_ai(full_url, details["title"], details["summary"], source_name, municipality, details["pub_date"], "web")

def parse_web_source(source_name, url, municipality):
    res = request_with_retry("GET", url, headers=HEADERS, timeout=30, verify=False)
    if not res:
        record_source_result(url, source_name, success=False)
        return
    record_source_result(url, source_name, success=True)
    res.encoding = res.apparent_encoding or 'utf-8'
    soup = BeautifulSoup(res.text, "html.parser")

    seen_urls, links_to_process = set(), []
    for link in soup.find_all("a", href=True):
        raw_href = link["href"].strip()
        if len(raw_href) < 2 or raw_href.startswith(("javascript:", "mailto:", "tel:")): continue
        full_url = urljoin(url, raw_href)
        if not full_url.startswith(("http://", "https://")): continue
        if "." not in (full_url.split('/')[2] if len(full_url.split('/')) > 2 else ""): continue
        if full_url in seen_urls or full_url.rstrip("/") == url.rstrip("/"): continue
        if any(x in full_url.lower() for x in ["/tag/", "/category/", "/search", "/contacts", "/about"]): continue
        
        seen_urls.add(full_url)
        links_to_process.append(full_url)

    stats.add("links_found", len(links_to_process))
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = [executor.submit(process_single_web_news, link, source_name, municipality) for link in links_to_process]
        for _ in as_completed(futures): pass

# --- 3. ПАРСИНГ ИСТОЧНИКОВ MAX (PLAYWRIGHT) ---
def parse_max_sources_batch(max_sources):
    if not max_sources: return
    logger.info("==================================================")
    logger.info("🚀 ЗАПУСК БРАУЗЕРНОГО ДВИЖКА ДЛЯ ИСТОЧНИКОВ MAX")
    logger.info("==================================================")

    # Весь блок (включая запуск браузера и ensure_max_ready) обёрнут в try —
    # иначе просроченная сессия MAX (ensure_max_ready кидает RuntimeError в
    # headless-режиме) роняла бы весь run() целиком: терялся бы уже готовый
    # экспорт WEB, а Telegram вообще не запускался. У Telegram такая защита
    # уже была (см. parse_telegram_sources_batch), у MAX — не было.
    try:
        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                MAX_PROFILE_DIR, headless=True, viewport={'width': 1440, 'height': 1000}, locale='ru-RU', timezone_id='Europe/Moscow'
            )
            try:
                context.add_init_script(SOCKET_HOOK)
                page = context.pages[0] if context.pages else context.new_page()
                ensure_max_ready(page, headless=True)
                page.wait_for_function('window.__maxNewsProtocol?.ready()', timeout=30000)

                collector = ApiCollector(MaxProtocol(page), delay_min=1.0, delay_max=2.0)
                end_date = EFFECTIVE_END_DATE

                for index, src in enumerate(max_sources, 1):
                    logger.info(f"\nПарсинг MAX [{index}/{len(max_sources)}]: {src['name']} ({src['muni']})")
                    stats.add("max_sources")
                    try:
                        # Создаем объект Source, который требует парсер MAX
                        from max_news.models import Source
                        max_source_obj = Source(name=src['name'], url=src['url'])

                        for batch in collector.batches(max_source_obj, EFFECTIVE_START_DATE, end_date):
                            for post in batch:
                                stats.add("links_found")
                                if is_url_processed(post.direct_url): continue
                                process_and_save_ai(
                                    url=post.direct_url, title=post.title, text=post.text,
                                    source_name=src['name'], municipality=src['muni'],
                                    pub_date=post.publication_date.strftime("%Y-%m-%d"), platform="max"
                                )
                        record_source_result(src['url'], src['name'], success=True)
                    except Exception as e:
                        stats.add("errors")
                        record_source_result(src['url'], src['name'], success=False)
                        logger.error(f"❌ Ошибка сбора MAX для {src['name']}: {e}")
            finally:
                context.close()
    except Exception as e:
        stats.add("errors")
        logger.error(
            f"❌ Не удалось запустить браузер для MAX: {e} Если требуется вход — "
            "выполните `python login_max.py`."
        )

# --- 3.1. ПАРСИНГ ИСТОЧНИКОВ TELEGRAM (PLAYWRIGHT, ТОТ ЖЕ ПРИНЦИП, ЧТО У MAX) ---
def parse_telegram_sources_batch(telegram_sources):
    if not telegram_sources: return
    logger.info("==================================================")
    logger.info("🚀 ЗАПУСК БРАУЗЕРНОГО ДВИЖКА ДЛЯ ИСТОЧНИКОВ TELEGRAM")
    logger.info("==================================================")

    end_date = EFFECTIVE_END_DATE
    try:
        with SocialBrowser("telegram", TELEGRAM_PROFILE_DIR, headless=True, logger=logger) as browser:
            for index, src in enumerate(telegram_sources, 1):
                logger.info(f"\nПарсинг Telegram [{index}/{len(telegram_sources)}]: {src['name']} ({src['muni']})")
                stats.add("telegram_sources")
                try:
                    for post in browser.collect(src['url'], EFFECTIVE_START_DATE, end_date):
                        stats.add("links_found")
                        if is_url_processed(post.direct_url): continue
                        process_and_save_ai(
                            url=post.direct_url, title=post.title, text=post.text,
                            source_name=src['name'], municipality=src['muni'],
                            pub_date=post.publication_date.strftime("%Y-%m-%d"), platform="telegram"
                        )
                    record_source_result(src['url'], src['name'], success=True)
                except AuthenticationRequired:
                    # Не вина конкретного источника — не засчитываем как провал источника.
                    stats.add("errors")
                    logger.error(
                        "❌ Требуется вход в Telegram: выполните `python login_telegram.py` "
                        "(один раз, вручную) и повторите сбор — остальные источники Telegram пропущены."
                    )
                    break
                except SocialBrowserError as e:
                    stats.add("errors")
                    record_source_result(src['url'], src['name'], success=False)
                    logger.error(f"❌ Ошибка сбора Telegram для {src['name']}: {e}")
    except Exception as e:
        stats.add("errors")
        logger.error(f"❌ Не удалось запустить браузер для Telegram: {e}")

# --- 4. ФИНАЛЬНАЯ СВОДКА И ЭКСПОРТ ---
def _print_boxed(title, lines):
    border = "═" * 40
    print(f"\n\033[96m╔{border}╗")
    print(f"║{title.center(40)}║")
    print(f"╠{border}╣")
    for line in lines: print(f"║ {line.ljust(39)}║")
    print(f"╚{border}╝\033[0m")

def print_summary():
    _print_boxed("ИТОГИ МЕГА-СБОРА НОВОСТЕЙ", [
        f"Обычных сайтов (WEB):  {stats.web_sources}",
        f"Источников MAX:        {stats.max_sources}",
        f"Источников Telegram:   {stats.telegram_sources}",
        f"В чёрном списке (пропущ.): {stats.blacklisted_skipped}",
        f"Всего новостей найдено:{stats.links_found}",
        f"Запросов к GigaChat:   {stats.gigachat_calls}",
        f"Отклонено нейросетью:  {stats.ai_rejected}",
        f"Ошибок (сеть/API):     {stats.errors}",
        f"Успешно сохранено:     {stats.saved}",
    ])
    print()

def platform_output_path(platform):
    """Путь для файла, отдельного по одному движку — то же имя, что у общего
    итогового файла, плюс суффикс (…_web.xlsx / …_max.xlsx / …_telegram.xlsx)."""
    return NEWS_OUTPUT_FILE.with_name(f"{NEWS_OUTPUT_FILE.stem}_{platform}{NEWS_OUTPUT_FILE.suffix}")

def export_db_to_excel(platform=None, output_path=None):
    """platform=None — как раньше, все новости в общий файл. Один из
    "web"/"max"/"telegram" — только новости этого движка (по колонке
    parsed_news.platform) в отдельный файл (platform_output_path() по
    умолчанию, если output_path не задан явно)."""
    target_path = output_path or (platform_output_path(platform) if platform else NEWS_OUTPUT_FILE)
    conn = sqlite3.connect(DB_FILE)
    if platform:
        df_db = pd.read_sql_query("SELECT * FROM parsed_news WHERE platform = ?", conn, params=(platform,))
    else:
        df_db = pd.read_sql_query("SELECT * FROM parsed_news", conn)
    conn.close()

    if df_db.empty:
        logger.info(f"ℹ️ Нет новостей для экспорта{f' ({platform})' if platform else ''} — файл не создаётся.")
        return

    try:
        wb = openpyxl.load_workbook(NEWS_TEMPLATE_FILE)
        sheet = wb["Новости"] if "Новости" in wb.sheetnames else wb.active
    except FileNotFoundError:
        logger.warning(
            f"⚠️ Шаблон {NEWS_TEMPLATE_FILE} не найден — создаю файл с голыми заголовками, "
            "без форматирования, инструкций и выпадающих списков из шаблона."
        )
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Новости"
        sheet.append(["Сборщик", "Дата сбора", "Основной муниципалитет", "Дополнительные муниципалитеты", "Дата публикации", "Дата события", "Источник", "Прямая ссылка", "Дополнительные ссылки", "Заголовок", "Категория", "Стадия события", "Краткое фактическое резюме", "Ключевые цифры и факты", "Упомянутые лица", "Организации", "ИНН из источника", "Объект/населённый пункт"])

    start_row = max([4] + [r for r in range(5, sheet.max_row + 1) if sheet.cell(row=r, column=8).value]) + 1
    excel_urls = set(str(sheet.cell(row=r, column=8).value).strip() for r in range(6, start_row) if sheet.cell(row=r, column=8).value)

    current_row = start_row
    for _, row in df_db.iterrows():
        url = str(row['url']).strip()
        if url in excel_urls: continue
        sheet.cell(row=current_row, column=1, value="Python parser")
        sheet.cell(row=current_row, column=2, value=datetime.now().strftime("%Y-%m-%d"))
        sheet.cell(row=current_row, column=3, value=row['municipality'])
        sheet.cell(row=current_row, column=5, value=row['pub_date'])
        sheet.cell(row=current_row, column=7, value=row['source'])
        sheet.cell(row=current_row, column=8, value=url)
        sheet.cell(row=current_row, column=10, value=row['title'])
        sheet.cell(row=current_row, column=11, value=row['category'])
        sheet.cell(row=current_row, column=12, value=row['stage'])
        sheet.cell(row=current_row, column=13, value=row['summary'])
        sheet.cell(row=current_row, column=14, value=row['facts'])
        sheet.cell(row=current_row, column=15, value=row['persons'])
        sheet.cell(row=current_row, column=16, value=row['organizations'])
        sheet.cell(row=current_row, column=17, value=row['inn'])
        sheet.cell(row=current_row, column=18, value=row['location'])
        current_row += 1

    try:
        wb.save(target_path)
    except PermissionError:
        # Частый случай — файл открыт в Excel на момент сохранения. Данные уже
        # надёжно лежат в БД (parsed_news), поэтому это не должно ронять весь
        # run() трейсбеком: просто пропускаем экспорт в этот раз, следующий
        # успешный запуск экспортирует все накопленные строки как обычно.
        logger.error(
            f"❌ Не удалось сохранить {target_path.name} — файл открыт в другой программе "
            "(например, в Excel). Закройте его и запустите сбор ещё раз: все новости уже "
            "сохранены в базе, экспорт просто повторится."
        )
        return
    logger.info(f"✅ Экспорт завершен. Файл: {target_path}")

ALL_STAGES = frozenset({"web", "max", "telegram"})

# --- ЗАПУСК ---
def run(stages=None):
    """stages: подмножество {"web", "max", "telegram"} — какие движки запускать
    в этот раз. None (по умолчанию) = все три, как раньше."""
    global EFFECTIVE_START_DATE, EFFECTIVE_END_DATE
    stages = ALL_STAGES if stages is None else frozenset(stages)
    stats.reset()
    init_db()
    if not os.path.exists(SOURCES_FILE):
        logger.error(f"Файл источников не найден: {SOURCES_FILE}")
        return
    if stages != ALL_STAGES:
        logger.info(f"▶️ Запуск только по этапам: {', '.join(sorted(stages))}")

    manual_window = get_manual_start_date() is not None or get_manual_end_date() is not None
    EFFECTIVE_START_DATE = compute_effective_start_date()
    EFFECTIVE_END_DATE = compute_effective_end_date()
    if manual_window:
        # Ручное окно действует один раз — дальше снова автоматический инкремент,
        # иначе оно бы навсегда «застряло» и следующие запуски не продвигались.
        set_manual_start_date(None)
        set_manual_end_date(None)
    logger.info(
        f"📅 Окно сбора: с {EFFECTIVE_START_DATE} по {EFFECTIVE_END_DATE} "
        f"(нижняя граница проекта — {START_DATE}){' [ручное окно, разовое]' if manual_window else ''}"
    )

    df_sources = pd.read_excel(SOURCES_FILE, sheet_name="Источники", header=4)
    web_sources = []
    max_sources = []
    telegram_sources = []
    blacklisted_urls = get_blacklisted_urls()

    # Распределяем источники по типу
    for idx, row in df_sources.iterrows():
        url, name, muni, status = [str(row.get(k, "")).strip() for k in ["Основная ссылка", "Название", "Муниципалитет", "Активность"]]
        if not url.startswith(("http://", "https://")) or (status and status.lower() not in ["активен", "nan", "none"]):
            continue

        if url in blacklisted_urls:
            stats.add("blacklisted_skipped")
            continue

        platform = str(row.get("Платформа", "")).strip().lower()
        # Если ссылка содержит max.ru или платформа указана как MAX — отправляем в браузерный парсер MAX
        if "max.ru" in url.lower() or platform == "max":
            max_sources.append({"name": name, "url": url, "muni": muni})
        # Telegram требует входа для полной истории так же, как MAX — тот же браузерный движок
        elif platform == "telegram" or "t.me" in url.lower() or "telegram.me" in url.lower():
            telegram_sources.append({"name": name, "url": url, "muni": muni})
        else:
            web_sources.append({"name": name, "url": url, "muni": muni})

    if stats.blacklisted_skipped:
        logger.warning(
            f"🚫 Пропущено источников из чёрного списка: {stats.blacklisted_skipped} "
            f"(после {SOURCE_FAILURE_THRESHOLD}+ провалов подряд). Список: --list-blacklist, сброс: --reset-blacklist"
        )

    # ЭТАП 1: Быстрый многопоточный парсинг обычных сайтов
    if "web" in stages:
        for index, src in enumerate(web_sources, 1):
            stats.add("web_sources")
            logger.info(f"\nПарсинг WEB [{index}/{len(web_sources)}]: {src['name']} ({src['muni']})")
            parse_web_source(src['name'], src['url'], src['muni'])

    # ЭТАП 2: Парсинг защищенных источников MAX через Playwright
    if "max" in stages:
        parse_max_sources_batch(max_sources)

    # ЭТАП 3: Парсинг защищенных источников Telegram через Playwright (тот же принцип входа)
    if "telegram" in stages:
        parse_telegram_sources_batch(telegram_sources)

    # ЭТАП 4: Финальный экспорт — собирали все три движка сразу, как раньше,
    # один общий файл; собирали не все — отдельный файл на каждый реально
    # запущенный в этот раз движок (…_web.xlsx / …_max.xlsx / …_telegram.xlsx).
    if stages == ALL_STAGES:
        export_db_to_excel()
    else:
        for p in sorted(stages):
            export_db_to_excel(platform=p)
    print_summary()
    if stages != ALL_STAGES:
        logger.info(
            "ℹ️ Запуск был частичным — дата последнего полного сбора для автоматического "
            "окна не обновлена, чтобы пропущенные сейчас этапы не потеряли часть периода."
        )
        return
    # Не date.today(): при разовом ручном окне с концом раньше сегодняшнего дня
    # автоматический режим следующего запуска должен продолжить именно с этой
    # даты, а не решить, что всё уже проверено по сегодня включительно.
    set_last_run_date(EFFECTIVE_END_DATE)

# --- 5. ИНТЕРАКТИВНОЕ МЕНЮ ---
def _read_choice(prompt):
    try:
        return input(prompt).strip()
    except (EOFError, KeyboardInterrupt):
        print("\nВыход.")
        sys.exit(0)

def show_blacklist_menu():
    while True:
        _print_boxed("ЧЁРНЫЙ СПИСОК ИСТОЧНИКОВ", [
            "1. Показать список",
            "2. Очистить весь список",
            "3. Очистить один источник",
            "4. Назад",
        ])
        choice = _read_choice("Выбор: ")
        if choice == "1":
            print_blacklist()
        elif choice == "2":
            confirm = _read_choice("Точно очистить весь чёрный список? (да/нет): ").lower()
            if confirm in ("да", "д", "y", "yes"):
                removed = reset_blacklist(None)
                print(f"♻️ Снято {removed} источник(ов).")
            else:
                print("Отменено.")
        elif choice == "3":
            target = _read_choice("Название или ссылка источника (частичное совпадение по названию): ")
            if target:
                removed = reset_blacklist(target)
                print(f"♻️ Снято {removed} источник(ов) по запросу «{target}».")
            else:
                print("Пустой ввод — отменено.")
        elif choice == "4":
            return
        else:
            print("Неверный выбор, введите число от 1 до 4.")

def show_paths_menu():
    global SOURCES_FILE, NEWS_TEMPLATE_FILE, NEWS_OUTPUT_FILE, DB_FILE
    while True:
        _print_boxed("ПУТИ К ФАЙЛАМ", [
            "1. Источники:",
            f"   {SOURCES_FILE}",
            "2. Шаблон новостей:",
            f"   {NEWS_TEMPLATE_FILE}",
            "3. Итоговый файл:",
            f"   {NEWS_OUTPUT_FILE}",
            "4. База данных:",
            f"   {DB_FILE}",
            "5. Сбросить к путям по умолчанию",
            "6. Назад",
        ])
        choice = _read_choice("Выбор: ")
        if choice in ("1", "2", "3", "4"):
            label = {"1": "источникам", "2": "шаблону новостей", "3": "итоговому файлу", "4": "базе данных"}[choice]
            new_path = _read_choice(f"Новый путь к {label} (Enter — оставить как есть): ")
            if not new_path:
                print("Оставлено без изменений.")
                continue
            path = Path(new_path)
            if choice in ("1", "2") and not path.exists():
                print(f"⚠️ Файл не найден по этому пути — сохраняю всё равно, проверьте перед сбором: {path}")
            if choice == "1": SOURCES_FILE = path
            elif choice == "2": NEWS_TEMPLATE_FILE = path
            elif choice == "3": NEWS_OUTPUT_FILE = path
            elif choice == "4":
                DB_FILE = path
                init_db()  # на новом пути таблиц ещё нет — создаём сразу
            save_paths_config()
            print(f"Сохранено: {path}")
        elif choice == "5":
            confirm = _read_choice("Сбросить все пути к значениям по умолчанию? (да/нет): ").lower()
            if confirm in ("да", "д", "y", "yes"):
                SOURCES_FILE = DEFAULT_SOURCES_FILE
                NEWS_TEMPLATE_FILE = DEFAULT_NEWS_TEMPLATE_FILE
                NEWS_OUTPUT_FILE = DEFAULT_NEWS_OUTPUT_FILE
                DB_FILE = DEFAULT_DB_FILE
                if PATHS_CONFIG_FILE.exists():
                    PATHS_CONFIG_FILE.unlink()
                init_db()
                print("Сброшено к путям по умолчанию.")
            else:
                print("Отменено.")
        elif choice == "6":
            return
        else:
            print("Неверный выбор, введите число от 1 до 6.")

def _parse_date_input(text):
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        print("Неверный формат даты, используйте ГГГГ-ММ-ДД (например, 2026-08-01).")
        return None

def show_date_window_menu():
    while True:
        manual_start = get_manual_start_date()
        manual_end = get_manual_end_date()
        preview_start = manual_start if manual_start is not None else compute_effective_start_date()
        preview_end = manual_end if manual_end is not None else date.today()
        mode = "ручной (разово, на следующий запуск)" if (manual_start or manual_end) else "автоматический"
        _print_boxed("ПЕРИОД ПОИСКА НОВОСТЕЙ", [
            f"Режим: {mode}",
            f"Применится: с {preview_start} по {preview_end}",
            "1. Задать дату начала",
            "2. Задать дату окончания",
            "3. Сбросить к автоматическому режиму",
            "4. Назад",
        ])
        choice = _read_choice("Выбор: ")
        if choice == "1":
            text = _read_choice("Дата начала (ГГГГ-ММ-ДД, Enter — отмена): ")
            if not text:
                print("Отменено.")
                continue
            d = _parse_date_input(text)
            if d is None:
                continue
            if d < START_DATE:
                print(f"⚠️ Дата раньше нижней границы проекта ({START_DATE}) — сохраняю как вы задали.")
            set_manual_start_date(d)
            print(f"Дата начала зафиксирована: {d}. Действует один раз, на ближайший запуск сбора.")
        elif choice == "2":
            text = _read_choice("Дата окончания (ГГГГ-ММ-ДД, Enter — отмена): ")
            if not text:
                print("Отменено.")
                continue
            d = _parse_date_input(text)
            if d is None:
                continue
            set_manual_end_date(d)
            print(f"Дата окончания зафиксирована: {d}. Действует один раз, на ближайший запуск сбора.")
        elif choice == "3":
            set_manual_start_date(None)
            set_manual_end_date(None)
            print("Возвращено в автоматический режим.")
        elif choice == "4":
            return
        else:
            print("Неверный выбор, введите число от 1 до 4.")

def show_collect_menu():
    while True:
        _print_boxed("ЗАПУСК СБОРА", [
            "1. Все источники (WEB + MAX + Telegram)",
            "2. Только WEB",
            "3. Только MAX",
            "4. Только Telegram",
            "5. Назад",
        ])
        choice = _read_choice("Выбор: ")
        if choice == "1":
            run()
        elif choice == "2":
            run(stages={"web"})
        elif choice == "3":
            run(stages={"max"})
        elif choice == "4":
            run(stages={"telegram"})
        elif choice == "5":
            return
        else:
            print("Неверный выбор, введите число от 1 до 5.")

def show_export_menu():
    while True:
        _print_boxed("ЭКСПОРТ ПО ИСТОЧНИКАМ", [
            "1. WEB — отдельный файл",
            "2. MAX — отдельный файл",
            "3. Telegram — отдельный файл",
            "4. Все три отдельными файлами",
            "5. Назад",
        ])
        choice = _read_choice("Выбор: ")
        if choice == "1":
            export_db_to_excel(platform="web")
        elif choice == "2":
            export_db_to_excel(platform="max")
        elif choice == "3":
            export_db_to_excel(platform="telegram")
        elif choice == "4":
            for p in ("web", "max", "telegram"):
                export_db_to_excel(platform=p)
        elif choice == "5":
            return
        else:
            print("Неверный выбор, введите число от 1 до 5.")

def show_main_menu():
    while True:
        _print_boxed("МЕГА-ПАРСЕР НОВОСТЕЙ", [
            "1. Запустить сбор новостей",
            "2. Чёрный список источников",
            "3. Пути к файлам",
            "4. Период поиска новостей",
            "5. Экспорт по источникам (WEB/MAX/Telegram отдельно)",
            "6. Выход",
        ])
        choice = _read_choice("Выбор: ")
        if choice == "1":
            show_collect_menu()
        elif choice == "2":
            show_blacklist_menu()
        elif choice == "3":
            show_paths_menu()
        elif choice == "4":
            show_date_window_menu()
        elif choice == "5":
            show_export_menu()
        elif choice == "6":
            print("Выход.")
            return
        else:
            print("Неверный выбор, введите число от 1 до 6.")

def parse_cli_args():
    parser = argparse.ArgumentParser(description="Мега-парсер новостей")
    parser.add_argument(
        "--run", action="store_true",
        help="Запустить сбор сразу, без интерактивного меню (для планировщика/автоматизации)",
    )
    parser.add_argument(
        "--reset-blacklist", nargs="?", const="__ALL__", default=None, metavar="ИСТОЧНИК",
        help="Сбросить чёрный список источников: без значения — полностью, "
             "со значением — по названию/ссылке (частичное совпадение по названию)",
    )
    parser.add_argument(
        "--list-blacklist", action="store_true",
        help="Показать источники в чёрном списке и выйти, без сбора новостей",
    )
    parser.add_argument(
        "--only", nargs="+", choices=sorted(ALL_STAGES), default=None, metavar="ЭТАП",
        help="Собирать только указанные движки: web, max и/или telegram (через пробел). "
             "По умолчанию — все три. Требует --run.",
    )
    parser.add_argument(
        "--export-only", nargs="+", choices=sorted(ALL_STAGES), default=None, metavar="ЭТАП",
        help="Не собирать ничего, только выгрузить уже накопленные в БД новости отдельными "
             "файлами по движкам (…_web.xlsx / …_max.xlsx / …_telegram.xlsx)",
    )
    parser.add_argument(
        "--sources", type=Path, default=None, metavar="ФАЙЛ",
        help=f"Путь к таблице источников (по умолчанию: {SOURCES_FILE.name})",
    )
    parser.add_argument(
        "--news-template", type=Path, default=None, metavar="ФАЙЛ",
        help=f"Путь к шаблону таблицы новостей (по умолчанию: {NEWS_TEMPLATE_FILE.name})",
    )
    parser.add_argument(
        "--output", type=Path, default=None, metavar="ФАЙЛ",
        help=f"Путь для итогового файла новостей (по умолчанию: {NEWS_OUTPUT_FILE.name})",
    )
    parser.add_argument(
        "--db", type=Path, default=None, metavar="ФАЙЛ",
        help=f"Путь к файлу базы данных (по умолчанию: {DB_FILE.name})",
    )
    return parser.parse_args()

if __name__ == "__main__":
    args = parse_cli_args()

    # Сохранённые через меню «Пути к файлам» пути — раньше CLI-флагов, чтобы
    # явный флаг всегда мог их перебить для разового/планового запуска.
    load_paths_config()

    # Переопределяем пути ДО init_db()/меню — они читаются по имени модуля-глобала
    # везде ниже, так что достаточно один раз переприсвоить здесь.
    if args.sources is not None:
        SOURCES_FILE = args.sources
        logger.info(f"📄 Источники: {SOURCES_FILE}")
    if args.news_template is not None:
        NEWS_TEMPLATE_FILE = args.news_template
        logger.info(f"📄 Шаблон новостей: {NEWS_TEMPLATE_FILE}")
    if args.output is not None:
        NEWS_OUTPUT_FILE = args.output
        logger.info(f"📄 Итоговый файл: {NEWS_OUTPUT_FILE}")
    if args.db is not None:
        DB_FILE = args.db
        logger.info(f"📄 База данных: {DB_FILE}")

    init_db()
    any_flag = (args.run or args.reset_blacklist is not None or args.list_blacklist
                or args.only is not None or args.export_only is not None)

    if not any_flag:
        # Без флагов — запуск вручную из консоли: показываем меню вместо мгновенного сбора.
        show_main_menu()
    else:
        # С флагами — скриптовый/плановый запуск: делаем ровно то, что попросили, без меню.
        if args.reset_blacklist is not None:
            target = None if args.reset_blacklist == "__ALL__" else args.reset_blacklist
            removed = reset_blacklist(target)
            logger.info(f"♻️ Чёрный список сброшен: снято {removed} источник(ов).")
        if args.list_blacklist:
            print_blacklist()
        if args.run:
            run(stages=set(args.only) if args.only else None)
        elif args.only:
            logger.error("❌ --only без --run ничего не запускает — добавьте --run.")
        if args.export_only:
            for p in args.export_only:
                export_db_to_excel(platform=p)