import os
import re
import json
import sqlite3
import pandas as pd
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from datetime import datetime, date
from pathlib import Path
import openpyxl
import urllib3
import uuid
import time
import warnings
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock

# --- НАСТРОЙКА ЛОГИРОВАНИЯ ---
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler("parser.log", encoding="utf-8"),
        logging.StreamHandler()
    ]
)

# Отключаем системные предупреждения
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
os.system("") 

# --- НАСТРОЙКИ ПУТЕЙ И API ---
BASE_DIR = Path(__file__).resolve().parent
SOURCES_FILE = BASE_DIR / "Мониторинг_Источники.xlsx"
NEWS_TEMPLATE_FILE = BASE_DIR / "Мониторинг_Новости.xlsx"
NEWS_OUTPUT_FILE = BASE_DIR / "Мониторинг_Новости_заполненный.xlsx"
DB_FILE = BASE_DIR / "news_database.db"

GIGACHAT_CREDENTIALS = "MDE5YmFjZDMtMzFhYi03N2Y4LTkzODQtZDkwNjAzZjgxMTMxOjE1MGJiYTcyLTRlYjMtNDNiZS05NjA3LTgwZDlmYzc4MTU1MA=="
START_DATE = date(2026, 7, 1)

# Снижаем нагрузку на сеть и процессор
MAX_WORKERS = 2 

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

class ParserStats:
    def __init__(self):
        self.lock = Lock()
        self.sources = 0
        self.links_found = 0
        self.ai_rejected = 0
        self.saved = 0
        self.errors = 0
        
    def add(self, metric, count=1):
        with self.lock:
            setattr(self, metric, getattr(self, metric) + count)

stats = ParserStats()
db_lock = Lock()

class ColoredConsoleFormatter(logging.Formatter):
    def format(self, record):
        msg = record.getMessage()
        time_str = self.formatTime(record, "%H:%M:%S")
        GRAY, RESET = "\033[90m", "\033[0m"
        GREEN, YELLOW, RED = "\033[92m", "\033[93m", "\033[91m"
        CYAN, MAGENTA, WHITE = "\033[96m", "\033[95m", "\033[97m"
        
        if "✅" in msg: color = GREEN
        elif "🗑️" in msg or "Пропуск" in msg: color = YELLOW
        elif "❌" in msg or "Ошибка" in msg: color = RED
        elif "🤖" in msg or "✨" in msg: color = CYAN
        elif "Парсинг" in msg: color = MAGENTA
        else: color = WHITE
        
        return f"{GRAY}[{time_str}]{RESET} {color}{msg}{RESET}"

logger = logging.getLogger()
logger.setLevel(logging.INFO)
file_handler = logging.FileHandler("parser.log", encoding="utf-8")
file_handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
logger.addHandler(file_handler)
console_handler = logging.StreamHandler()
console_handler.setFormatter(ColoredConsoleFormatter())
logger.addHandler(console_handler)

# --- 1. ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ И СЕТЬ ---
def clean_text(text):
    if not text: return ""
    return re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F]', '', str(text))[:32000].strip()

def extract_date_obj(text):
    if not text: return None
    text = str(text).strip().lower()
    match_iso = re.search(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', text)
    if match_iso:
        y, m, d = map(int, match_iso.groups())
        if 2000 <= y <= 2030 and 1 <= m <= 12 and 1 <= d <= 31:
            try: return date(y, m, d)
            except ValueError: pass
    match_ru = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', text)
    if match_ru:
        d, m, y = map(int, match_ru.groups())
        if 2000 <= y <= 2030 and 1 <= m <= 12 and 1 <= d <= 31:
            try: return date(y, m, d)
            except ValueError: pass
    months = {'янв':1, 'фев':2, 'мар':3, 'апр':4, 'мая':5, 'май':5, 'июн':6, 'июл':7, 'авг':8, 'сен':9, 'окт':10, 'ноя':11, 'дек':12}
    match_words = re.search(r'(\d{1,2})\s+([а-яя]+)\s+(\d{4})', text)
    if match_words:
        d, m_word, y = match_words.groups()
        for key, val in months.items():
            if m_word.startswith(key):
                try: return date(int(y), val, int(d))
                except ValueError: pass
    return None

def request_with_retry(method, url, max_retries=3, **kwargs):
    delay = 2
    for attempt in range(max_retries):
        try:
            res = requests.request(method, url, **kwargs)
            res.raise_for_status()
            return res
        except requests.exceptions.RequestException as e:
            if attempt == max_retries - 1:
                stats.add("errors")
                logging.error(f"❌ Сбой сети {url[:50]}...: {e}")
                return None
            time.sleep(delay)
            delay *= 2

# --- 2. БАЗА ДАННЫХ ---
def init_db():
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute('''CREATE TABLE IF NOT EXISTS parsed_news (
            url TEXT PRIMARY KEY, title TEXT, source TEXT, municipality TEXT, pub_date TEXT, category TEXT, stage TEXT, summary TEXT, facts TEXT, persons TEXT, organizations TEXT, inn TEXT, location TEXT, created_at TEXT)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS rejected_urls (
            url TEXT PRIMARY KEY, reason TEXT, created_at TEXT)''')
        conn.commit()
        conn.close()

def is_url_processed(url):
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        if cursor.execute("SELECT 1 FROM parsed_news WHERE url = ?", (url,)).fetchone() or cursor.execute("SELECT 1 FROM rejected_urls WHERE url = ?", (url,)).fetchone():
            conn.close()
            return True
        conn.close()
        return False

def save_rejected_url(url, reason):
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        try:
            conn.execute("INSERT INTO rejected_urls (url, reason, created_at) VALUES (?, ?, ?)", (url, reason, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
            conn.commit()
        except sqlite3.IntegrityError: pass
        conn.close()

def save_item_to_db(item):
    def safe_str(val): return ", ".join(str(x) for x in val) if isinstance(val, list) else (str(val) if val is not None else "")
    with db_lock:
        conn = sqlite3.connect(DB_FILE)
        success = False
        try:
            conn.execute('''INSERT INTO parsed_news (url, title, source, municipality, pub_date, category, stage, summary, facts, persons, organizations, inn, location, created_at) 
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', 
                         (safe_str(item["url"]), safe_str(item["title"]), safe_str(item["source"]), safe_str(item["municipality"]), safe_str(item["pub_date"]), safe_str(item.get("category", "")), safe_str(item.get("stage", "")), safe_str(item.get("summary", "")), safe_str(item.get("facts", "")), safe_str(item.get("persons", "")), safe_str(item.get("organizations", "")), safe_str(item.get("inn", "")), safe_str(item.get("location", "")), datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
            conn.commit()
            success = True
        except sqlite3.IntegrityError: pass
        conn.close()
        return success

# --- 3. GIGACHAT API (ДВА ПРОСТЫХ ШАГА) ---
GIGACHAT_TOKEN = None
GIGACHAT_TOKEN_EXPIRES = 0

def get_gigachat_token():
    global GIGACHAT_TOKEN, GIGACHAT_TOKEN_EXPIRES
    if GIGACHAT_TOKEN and time.time() < GIGACHAT_TOKEN_EXPIRES: return GIGACHAT_TOKEN
    url = "https://ngw.devices.sberbank.ru:9443/api/v2/oauth"
    headers = {"Content-Type": "application/x-www-form-urlencoded", "Accept": "application/json", "RqUID": str(uuid.uuid4()), "Authorization": f"Basic {GIGACHAT_CREDENTIALS}"}
    res = request_with_retry("POST", url, data={"scope": "GIGACHAT_API_PERS"}, headers=headers, verify=False, timeout=15)
    if res:
        data = res.json()
        GIGACHAT_TOKEN = data["access_token"]
        GIGACHAT_TOKEN_EXPIRES = time.time() + (data["expires_at"] / 1000) - 60
        return GIGACHAT_TOKEN
    return None

def call_gigachat_api(system_prompt, user_message):
    token = get_gigachat_token()
    if not token: return None
    url = "https://gigachat.devices.sberbank.ru/api/v1/chat/completions"
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {token}"}
    payload = {
        "model": "GigaChat", 
        "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_message}], 
        "temperature": 0.1, 
        "max_tokens": 1000
    }
    res = request_with_retry("POST", url, headers=headers, json=payload, verify=False, timeout=30)
    if res:
        try: return json.loads(res.json()["choices"][0]["message"]["content"].replace("```json", "").replace("```", "").strip())
        except Exception: stats.add("errors")
    return None

# ШАГ 1: Простой фильтр
def is_news_relevant(title, text, municipality):
    system_prompt = f"""Ты строгий фильтр новостей. Муниципалитет: {municipality}.
    Правила:
    1. Если новость про концерты, праздники, афиши, советы или рекламу -> is_relevant: false
    2. Если новость про ремонты, стройку, бюджет, ДТП, экономику, решения властей -> is_relevant: true

    Верни ТОЛЬКО JSON:
    {{
        "is_relevant": true или false,
        "reason": "краткая причина"
    }}"""
    return call_gigachat_api(system_prompt, f"Заголовок: {title}\nТекст: {text}")

# ШАГ 2: Простой извлекатель
def extract_news_details(title, text, municipality):
    system_prompt = f"""Ты помощник по извлечению фактов. Муниципалитет: {municipality}.
    Заполни поля по тексту новости. 

    Категория (ВЫБЕРИ ОДНУ ИЗ): Власть и кадровые решения, Бюджет и муниципальные финансы, Закупки и контракты, Строительство и инфраструктура, ЖКХ и благоустройство, Экономика, бизнес и занятость, Социальная сфера, Образование, Здравоохранение, Происшествия и безопасность, Экология, Обращения и проблемы жителей, Культура, спорт и туризм, Банки и финансовые сервисы, Другое.
    Стадия (ВЫБЕРИ ОДНУ ИЗ): Анонс, Решение принято, Финансирование выделено, Закупка объявлена, Работы начаты, В процессе, Завершено / открыто, Перенесено / приостановлено, Проблема / происшествие, Результат.

    Верни ТОЛЬКО JSON:
    {{
        "category": "категория",
        "stage": "стадия",
        "summary": "1-2 предложения сути",
        "facts": "цифры, деньги, сроки",
        "persons": "ФИО",
        "organizations": "ведомства, компании",
        "inn": "ИНН (если есть)",
        "location": "улица, объект"
    }}"""
    return call_gigachat_api(system_prompt, f"Заголовок: {title}\nТекст: {text}")

# --- 4. ПАРСИНГ ---
def parse_article_details(url):
    # Увеличен таймаут до 30 для тяжелых сайтов
    res = request_with_retry("GET", url, headers=HEADERS, timeout=30, verify=False)
    if not res or "text/html" not in res.headers.get("Content-Type", "").lower(): return None
    res.encoding = res.apparent_encoding or 'utf-8'
    soup = BeautifulSoup(res.text, "html.parser")
    
    pub_date_obj = None
    for script in soup.find_all("script", type="application/ld+json"):
        if not script.string: continue
        try:
            data = json.loads(script.string.strip())
            for item in (data if isinstance(data, list) else [data]):
                if isinstance(item, dict):
                    for df in ["datePublished", "dateCreated", "uploadDate"]:
                        if df in item and isinstance(item[df], str):
                            pub_date_obj = extract_date_obj(item[df])
                            if pub_date_obj: break
                if pub_date_obj: break
        except Exception: continue

    if not pub_date_obj:
        for meta_name in ["article:published_time", "pubdate", "og:pubdate", "dc.date.issued", "parsely-pub-date", "date"]:
            tag = soup.find("meta", attrs={"property": meta_name}) or soup.find("meta", attrs={"name": meta_name})
            if tag and tag.get("content"):
                pub_date_obj = extract_date_obj(tag["content"])
                if pub_date_obj: break
    if not pub_date_obj:
        time_tag = soup.find("time")
        if time_tag: pub_date_obj = extract_date_obj(time_tag.get("datetime") or time_tag.get_text())

    if not pub_date_obj or not (START_DATE <= pub_date_obj <= date.today()): return None

    title = soup.find("h1").get_text(" ", strip=True) if soup.find("h1") else ""
    if not title or len(title) < 10:
        og_title = soup.find("meta", attrs={"property": "og:title"}) or soup.find("meta", attrs={"name": "twitter:title"})
        title = og_title["content"].strip() if og_title and og_title.get("content") else title
    if not title or len(title) < 10:
        title = soup.title.string.split("—")[0].split("-")[0].strip() if soup.title and soup.title.string else title

    paragraphs = [p.get_text(" ", strip=True) for p in soup.find_all("p") if len(p.get_text(" ", strip=True)) > 40]
    summary = " ".join(paragraphs[:3]) if paragraphs else ""
    if not summary:
        meta_desc = soup.find("meta", attrs={"name": "description"}) or soup.find("meta", attrs={"property": "og:description"})
        summary = meta_desc["content"].strip() if meta_desc and meta_desc.get("content") else summary

    return {"title": clean_text(title or summary[:100]+"..."), "pub_date": pub_date_obj.strftime("%Y-%m-%d"), "summary": clean_text(summary or title)}

def process_single_news(full_url, source_name, municipality):
    if is_url_processed(full_url): return None
    details = parse_article_details(full_url)
    if not details: return None

    logging.info(f"🤖 Проверка: {details['title'][:55]}...")
    filter_data = is_news_relevant(details["title"], details["summary"], municipality)
    
    if not filter_data: 
        return None
        
    if not filter_data.get("is_relevant"):
        stats.add("ai_rejected")
        reason = filter_data.get("reason", "Мусор")
        save_rejected_url(full_url, reason)
        logging.info(f"🗑️ [ОТКЛОНЕНО]: {reason}")
        return None

    logging.info(f"✨ Извлечение данных: {details['title'][:30]}...")
    ai_data = extract_news_details(details["title"], details["summary"], municipality)
    if not ai_data: ai_data = {}

    if details.get("title"):
        item = {"url": full_url, "title": details["title"], "source": source_name, "municipality": municipality, "pub_date": details["pub_date"], **{k: ai_data.get(k, "") for k in ["category", "stage", "summary", "facts", "persons", "organizations", "inn", "location"]}}
        if save_item_to_db(item):
            stats.add("saved")
            logging.info(f"✅ [СОХРАНЕНО | {item.get('category', 'Б/К')}] {item['title'][:45]}...")
            return item
    return None

def parse_source_news(source_name, url, municipality):
    # Увеличен таймаут
    res = request_with_retry("GET", url, headers=HEADERS, timeout=30, verify=False)
    if not res: return []
    res.encoding = res.apparent_encoding or 'utf-8'
    soup = BeautifulSoup(res.text, "html.parser")

    seen_urls, links_to_process = set(), []
    for link in soup.find_all("a", href=True):
        full_url = urljoin(url, link["href"].strip())
        if any(bad in full_url.lower() for bad in ["max.ru", "uggs.rkomi.ru"]) or not full_url.startswith(("http://", "https://")) or full_url in seen_urls: continue
        if full_url.rstrip("/") == url.rstrip("/") or any(x in full_url.lower() for x in ["/tag/", "/category/", "/search", "/contacts", "/about"]): continue
        seen_urls.add(full_url)
        links_to_process.append(full_url)

    stats.add("links_found", len(links_to_process))
    collected_items = []
    
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(process_single_news, link, source_name, municipality): link for link in links_to_process}
        for future in as_completed(futures):
            if result := future.result(): collected_items.append(result)

    return collected_items

# --- 5. ФИНАЛЬНАЯ СВОДКА И ЭКСПОРТ ---
def print_summary():
    lines = [
        f"Источников обработано: {stats.sources}",
        f"Всего ссылок найдено:  {stats.links_found}",
        f"Отклонено нейросетью:  {stats.ai_rejected}",
        f"Ошибок (сеть/API):     {stats.errors}",
        f"Успешно сохранено:     {stats.saved}"
    ]
    border = "═" * 40
    print(f"\n\033[96m╔{border}╗")
    print(f"║{'ИТОГИ СБОРА НОВОСТЕЙ'.center(40)}║")
    print(f"╠{border}╣")
    for line in lines: print(f"║ {line.ljust(39)}║")
    print(f"╚{border}╝\033[0m\n")

def export_db_to_excel():
    conn = sqlite3.connect(DB_FILE)
    df_db = pd.read_sql_query("SELECT * FROM parsed_news", conn)
    conn.close()

    if df_db.empty: return

    try:
        wb = openpyxl.load_workbook(NEWS_TEMPLATE_FILE)
        sheet = wb["Новости"] if "Новости" in wb.sheetnames else wb.active
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Новости"
        sheet.append(["Сборщик", "Дата сбора", "Основной муниципалитет", "Дополнительные муниципалитеты", "Дата публикации", "Дата события", "Источник", "Прямая ссылка", "Дополнительные ссылки", "Заголовок", "Категория", "Стадия события", "Краткое фактическое резюме", "Ключевые цифры и факты", "Упомянутые лица", "Организации", "ИНН из источника", "Объект/населённый пункт"])

    start_row = max([4] + [r for r in range(5, sheet.max_row + 1) if sheet.cell(row=r, column=8).value]) + 1
    excel_urls = set(str(sheet.cell(row=r, column=8).value).strip() for r in range(6, start_row) if sheet.cell(row=r, column=8).value)

    current_row = start_row
    for _, row in df_db.iterrows():
        url = str(row['url']).strip()
        if url in excel_urls: continue
        sheet.cell(row=current_row, column=1, value="Python parser")
        sheet.cell(row=current_row, column=2, value=datetime.now().strftime("%Y-%m-%d"))
        sheet.cell(row=current_row, column=3, value=clean_text(row['municipality']))
        sheet.cell(row=current_row, column=5, value=clean_text(row['pub_date']))
        sheet.cell(row=current_row, column=7, value=clean_text(row['source']))
        sheet.cell(row=current_row, column=8, value=clean_text(url))
        sheet.cell(row=current_row, column=10, value=clean_text(row['title']))
        sheet.cell(row=current_row, column=11, value=clean_text(row['category']))
        sheet.cell(row=current_row, column=12, value=clean_text(row['stage']))
        sheet.cell(row=current_row, column=13, value=clean_text(row['summary']))
        sheet.cell(row=current_row, column=14, value=clean_text(row['facts']))
        sheet.cell(row=current_row, column=15, value=clean_text(row['persons']))
        sheet.cell(row=current_row, column=16, value=clean_text(row['organizations']))
        sheet.cell(row=current_row, column=17, value=clean_text(row['inn']))
        sheet.cell(row=current_row, column=18, value=clean_text(row['location']))
        current_row += 1

    wb.save(NEWS_OUTPUT_FILE)
    logging.info(f"✅ Экспорт завершен. Файл: {NEWS_OUTPUT_FILE}")

# --- ЗАПУСК ---
def run():
    init_db()
    if not os.path.exists(SOURCES_FILE):
        logging.error(f"Файл источников не найден: {SOURCES_FILE}")
        return

    df_sources = pd.read_excel(SOURCES_FILE, sheet_name="Источники", header=4)
    
    for idx, row in df_sources.iterrows():
        url, name, muni, status = [str(row.get(k, "")).strip() for k in ["Основная ссылка", "Название", "Муниципалитет", "Активность"]]
        if not url.startswith(("http://", "https://")) or (status and status.lower() not in ["активен", "nan", "none"]):
            continue
        if any(bad in url.lower() for bad in ["max.ru", "uggs.rkomi.ru"]):
            logging.info(f"🟡 Пропуск источника: {name}")
            continue

        stats.add("sources")
        logging.info(f"==================================================")
        logging.info(f"Парсинг [{stats.sources}]: {name} ({muni})")
        logging.info(f"==================================================")
        
        parse_source_news(name, url, muni)

    export_db_to_excel()
    print_summary() 

if __name__ == "__main__":
    run()